from __future__ import annotations

import io

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .base import BaseExporter

HEADERS = ("Name", "Identifier", "Type", "Purpose", "Date", "Time In", "Time Out", "Duration")
WIDTHS = (24, 16, 12, 32, 12, 20, 20, 12)


class XlsxExporter(BaseExporter):
    key = "xlsx"
    label = "Excel"
    extension = "xlsx"
    content_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    def build(self, visits, generated_at) -> bytes:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Visits"

        sheet.append(HEADERS)
        header_fill = PatternFill("solid", fgColor="5C7CFA")
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(vertical="center")

        for visit in visits:
            sheet.append(
                (
                    visit.visitor.name,
                    visit.visitor.identifier,
                    visit.visitor.type_label,
                    visit.display_purpose,
                    visit.visit_date.isoformat(),
                    timezone.localtime(visit.time_in).strftime("%Y-%m-%d %I:%M %p"),
                    timezone.localtime(visit.time_out).strftime("%Y-%m-%d %I:%M %p")
                    if visit.time_out
                    else "",
                    visit.duration_label,
                )
            )

        for index, width in enumerate(WIDTHS, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = width

        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_view.showGridLines = False

        buffer = io.BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()
